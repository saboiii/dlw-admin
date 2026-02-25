from io import BytesIO
from pathlib import Path
from collections import Counter
from openpyxl import load_workbook

path = Path("/Volumes/XtremeLOAD/Projects/dlw-admin/lib/dlw_participants.csv")
blob = path.read_bytes()
wb = load_workbook(BytesIO(blob), read_only=True, data_only=True)
ws = wb.active

headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
course_idx = headers.index("Course")

courses = []
for row in ws.iter_rows(min_row=2, values_only=True):
    value = row[course_idx]
    if value is not None:
        courses.append(str(value).strip())

unique = list(dict.fromkeys(courses))

print("total courses:", len(courses))
print("unique courses:", len(unique))
print("\nSample unique values (first 50):")
for item in unique[:50]:
    print("-", item)

print("\nMost common values (top 30):")
for value, count in Counter(courses).most_common(30):
    print(f"- {value} ({count})")
